"""Parse the official PLR display-name workbook."""

import io
import re

from openpyxl import load_workbook

WORKSHEET_NAME = "LOR_2023_PLR"
HEADER = ("PLR_ID", "PLR_Name")


def parse_plr_names(content: bytes) -> dict[str, str]:
    """Return {plr_id: display name}. Raises ValueError on any irregularity."""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if WORKSHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"PLR workbook is missing worksheet {WORKSHEET_NAME}")
        rows = list(workbook[WORKSHEET_NAME].iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows or tuple(rows[0][:2]) != HEADER:
        raise ValueError("PLR worksheet must begin with PLR_ID and PLR_Name")

    names: dict[str, str] = {}
    for raw_id, raw_name in (row[:2] for row in rows[1:]):
        plr_id = "" if raw_id is None else str(raw_id).strip()
        name = "" if raw_name is None else str(raw_name).strip()
        if not re.fullmatch(r"[0-9]{8}", plr_id):
            raise ValueError(f"PLR identifier must contain exactly eight digits: {plr_id!r}")
        if not name:
            raise ValueError(f"PLR {plr_id} has an empty display name")
        if plr_id in names:
            raise ValueError(f"PLR name directory contains duplicate identifier {plr_id}")
        names[plr_id] = name
    return names
